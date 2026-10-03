"""Extração de texto: funções puras (bytes -> texto), sem rede e sem banco."""

import io
import logging
import re
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone

import docx
from pypdf import PdfReader

# pypdf avisa muito em PDF mal formado; o erro real sobe como exceção.
logging.getLogger("pypdf").setLevel(logging.ERROR)

MAX_TEXT_BYTES = 200_000
# PDF abaixo disso é tratado como sem texto (escaneado, depende de OCR).
MIN_TEXT_CHARS = 100

_LOCAL_TZ = timezone(timedelta(hours=-3))
_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
# Ligaduras que o pypdf devolve como um caractere só ("Identiﬁcação", "NormaƟva").
# Ɵ e Ʃ são como o Word/Calibri codifica "ti" e "tt" ao exportar PDF.
_LIGATURES = str.maketrans({
    "ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl", "ﬃ": "ffi", "ﬄ": "ffl", "ﬅ": "ft", "ﬆ": "st",
    "Ɵ": "ti", "Ʃ": "tt",
})


@dataclass
class TextResult:
    text: str
    truncated: bool = False
    doc_created_at: datetime | None = None
    pages: int | None = None

    @property
    def is_textless(self) -> bool:
        # Só PDF (pages preenchido) usa o mínimo: pouco texto ali indica página escaneada.
        # TXT, DOCX e Google Docs curtos são texto de verdade; só vazio conta.
        chars = len(self.text.strip())
        return chars < MIN_TEXT_CHARS if self.pages is not None else chars == 0


def clean_text(text: str) -> str:
    """Desfaz ligaduras, remove caracteres de controle (o Postgres recusa \\x00) e espaços redundantes."""
    text = _CONTROL.sub(" ", text.translate(_LIGATURES))
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\s*\n\s*\n\s*", "\n\n", text)
    return text.strip()


def clamp_text(text: str, max_bytes: int = MAX_TEXT_BYTES) -> tuple[str, bool]:
    """Corta o texto em max_bytes de UTF-8 sem quebrar caractere."""
    encoded = text.encode("utf-8")
    if len(encoded) <= max_bytes:
        return text, False
    return encoded[:max_bytes].decode("utf-8", errors="ignore"), True


def _aware(dt: datetime | None) -> datetime | None:
    if dt is None:
        return None
    return dt if dt.tzinfo else dt.replace(tzinfo=_LOCAL_TZ)


def _result(raw: str, max_bytes: int, **extra) -> TextResult:
    text, truncated = clamp_text(clean_text(raw), max_bytes)
    return TextResult(text=text, truncated=truncated, **extra)


def extract_pdf(data: bytes, max_bytes: int = MAX_TEXT_BYTES) -> TextResult:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("PDF protegido por senha")

    parts, size = [], 0
    for page in reader.pages:
        try:
            chunk = page.extract_text() or ""
        except Exception:  # uma página ruim não derruba o documento
            continue
        parts.append(chunk)
        size += len(chunk.encode("utf-8"))
        if size > max_bytes * 2:  # já passou do limite; não processa o resto do livro
            break

    try:
        created = reader.metadata.creation_date if reader.metadata else None
    except Exception:  # data em formato inválido no metadado
        created = None
    return _result("\n\n".join(parts), max_bytes, doc_created_at=_aware(created), pages=len(reader.pages))


def extract_docx(data: bytes, max_bytes: int = MAX_TEXT_BYTES) -> TextResult:
    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    created = document.core_properties.created
    if created and not created.tzinfo:
        created = created.replace(tzinfo=timezone.utc)  # core.xml grava em UTC
    return _result("\n".join(parts), max_bytes, doc_created_at=created)


def decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def extract_txt(data: bytes, max_bytes: int = MAX_TEXT_BYTES) -> TextResult:
    return _result(decode_text(data), max_bytes)


def from_plain_text(text: str, max_bytes: int = MAX_TEXT_BYTES) -> TextResult:
    """Texto que já chega pronto (export de Google Docs)."""
    return _result(text, max_bytes)
